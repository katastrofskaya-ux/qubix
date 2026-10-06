#!/usr/bin/env python3
"""Отчётность по всем живым клиентам — таблица Excel для ручного заполнения.

Автоматические колонки — из админки (ступень, канал, подарок до, сервер на
связи) и из карточек SALES (последнее сообщение клиента боту, последнее
сообщение бота). Ручные колонки — пустые: писала лично, где, что ответил,
следующий шаг и дата.

    python3 scripts/klienty.py                 # → data/klienty-YYYY-MM-DD.xlsx

Нужны ADMIN_QUBIX_COOKIE и YOUTRACK_API_TOKEN. Только чтение. В файле личные
данные клиентов (почта, Telegram) — в репозиторий не идёт (data/ в .gitignore).
"""
import datetime
import json
import os
import re
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import touch  # noqa: E402  (фильтры живых клиентов и запрос к админке)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

YT = "https://team.qubix.capital"
ADMIN_CARD = "https://admin.qubix.pro/customers/"
NOW = datetime.datetime.now(datetime.timezone.utc)
TEST_PAYMENTS = {46}  # №46 — оплата $49 02.07, тестовая; реальная оплата одна, июль (владелец 03.10)

CHANNEL = {"ADHUNT": "adhunt", "ADHUNT1": "adhunt", "AFF_INSIDE": ".Aff Inside", "AFFINSIDE": ".Aff Inside",
           "CPAGRAM": "CPAGRAM", "CPAGRAM1": "CPAGRAM", "FBKILLA": "FB-killa", "FBKILLA_SITE": "FB-killa",
           "IGAMINGNEWS": "iGaming News", "IGAMING_NEWS": "iGaming News", "PACAN": "Pacan",
           "PARTNERKIN_SITE": "Partnerkin", "PARTNEROFF_SITE": "Partneroff", "CPARIP": "CPA.RIP",
           "KHOMENOK": "Хоменок", "TGQUBIX": "наш канал"}

STAGES = ["Оплатил", "Сервер стоит", "Лицензия без сервера", "Только регистрация"]
MANUAL = ["Писала лично (дата)", "Где (TG / почта / звонок)", "Что ответил", "Следующий шаг", "Дата шага"]


def yt(path, **params):
    args = ["curl", "-sS", "-G", "-H", f"Authorization: Bearer {os.environ['YOUTRACK_API_TOKEN']}",
            "-H", "Accept: application/json"]
    for k, v in params.items():
        args += ["--data-urlencode", f"{k}={v}"]
    r = subprocess.run(args + [f"{YT}/api{path}"], capture_output=True, text=True, timeout=120)
    return json.loads(r.stdout)


def d(value):
    return touch.dt(value) if value else None


def fmt(x):
    return x.strftime("%d.%m") if x else ""


def main():
    customers = [c for c in touch.load() if not c.get("is_internal")]
    cards = yt("/issues", query="project:SALES", **{"$top": "3000", "fields": "idReadable,summary,description"})
    blobs = [(i["idReadable"], (i["summary"] or "") + " " + (i["description"] or "")) for i in cards]

    rows = []
    for c in customers:
        full = touch.get(f"/customers/{c['id']}")
        lic = full.get("licenses") or []
        seen = max((d(l.get("last_seen")) for l in lic if l.get("last_seen")), default=None)
        gift = max((d(l.get("paid_until")) for l in lic if l.get("paid_until")), default=None) or d(c.get("paid_until"))
        real_pay = [p for p in (full.get("payments") or [])
                    if p.get("status") == "paid" and (p.get("amount_usd") or 0) > 0 and c["number"] not in TEST_PAYMENTS]
        if real_pay:
            stage = "Оплатил"
        elif seen:
            stage = "Сервер стоит"
        elif lic:
            stage = "Лицензия без сервера"
        else:
            stage = "Только регистрация"

        n, email, tg = c["number"], (c.get("email") or "").lower(), str(c.get("telegram_chat_id") or "")
        card = next((i for i, b in blobs if re.search(rf"#{n}\b", b) or (email and email in b.lower())
                     or (tg and tg in b)), None)
        client_said = bot_said = None
        if card:
            for m in yt(f"/issues/{card}/comments", **{"$top": "5000", "fields": "created,text"}):
                t = datetime.datetime.fromtimestamp(m["created"] / 1000, datetime.timezone.utc)
                txt = m.get("text") or ""
                if txt.startswith("👤"):
                    client_said = t
                elif txt.startswith("🤖"):
                    bot_said = t

        code = c["_code"] if c["_code"] != "—" else ""
        rows.append({
            "stage": stage, "gift": gift, "seen": seen,
            "cells": [
                n, stage, fmt(d(c.get("created_at"))),
                CHANNEL.get(code, f"код {code}" if code else "без кода"),
                f"tg://user?id={tg}" if tg else "нет, только почта", c.get("email") or "",
                ("кончился " if gift and gift < NOW else "") + fmt(gift) if gift else "",
                fmt(seen), fmt(client_said), fmt(bot_said),
                f"{YT}/issue/{card}" if card else "нет карточки", ADMIN_CARD + c["id"],
            ],
        })

    # Порядок: ступень, внутри — у кого подарок кончается раньше (идущие сначала), потом свежий сигнал сервера.
    far = datetime.datetime.max.replace(tzinfo=datetime.timezone.utc)
    rows.sort(key=lambda r: (STAGES.index(r["stage"]),
                             (r["gift"] or far) < NOW,
                             r["gift"] if r["gift"] and r["gift"] >= NOW else far,
                             -(r["seen"].timestamp() if r["seen"] else 0)))

    head = ["№", "Ступень", "Пришёл", "Канал", "Telegram", "Почта", "Подарок до", "Сервер на связи",
            "Сам писал в бота", "Бот писал", "Карточка SALES", "Админка"] + MANUAL
    wb = Workbook()
    ws = wb.active
    ws.title = "Клиенты"
    ws.append(head)
    for r in rows:
        ws.append(r["cells"] + [""] * len(MANUAL))
    bold, manual_fill = Font(bold=True), PatternFill("solid", fgColor="FFF4CC")
    for i, h in enumerate(head, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if h in MANUAL:
            for row in range(1, len(rows) + 2):
                ws.cell(row=row, column=i).fill = manual_fill
        ws.column_dimensions[get_column_letter(i)].width = {"Почта": 26, "Что ответил": 40, "Следующий шаг": 30,
                                                             "Карточка SALES": 14, "Админка": 14,
                                                             "Telegram": 14}.get(h, 12)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    note = wb.create_sheet("Как читать")
    for line in [
        f"Срез {NOW:%d.%m.%Y %H:%M} UTC. Источники: админка admin.qubix.pro (живые клиенты, без is_internal, "
        "тестовых и своих), карточки SALES в YouTrack.",
        "Ступень: Оплатил — платёж с суммой больше нуля (№46 — тестовый, не считается); Сервер стоит — сервер "
        "хоть раз отчитался; Лицензия без сервера — подарок взят, сервер не отчитывался; Только регистрация.",
        "Подарок до — дата окончания лицензии. Сервер на связи — когда сервер последний раз отчитался.",
        "Сам писал в бота / Бот писал — последняя дата в карточке SALES.",
        "Жёлтые колонки — ручные: личная переписка в Telegram и почте в трекер не попадает.",
    ]:
        note.append([line])
    note.column_dimensions["A"].width = 120

    out = Path(__file__).parent.parent / "data" / f"klienty-{NOW:%Y-%m-%d}.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    counts = {s: sum(r["stage"] == s for r in rows) for s in STAGES}
    print(f"{out} — {len(rows)} клиентов: " + ", ".join(f"{s} {n}" for s, n in counts.items()))


if __name__ == "__main__":
    main()
